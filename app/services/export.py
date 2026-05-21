"""
Excel 导出服务 — 保留 openpyxl 美化样式
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

from app.core.config import settings
from app.core.exceptions import ExportError
from app.core.logging import get_logger
from app.services.screening import get_screened_stocks

logger = get_logger("export_service")


def export_screened_to_excel(trade_date: date | None = None) -> Path:
    """
    将筛选结果导出为 Excel 文件

    Args:
        trade_date: 交易日，None = 最近

    Returns:
        生成文件的 Path
    """
    try:
        import openpyxl
        from openpyxl.styles import (
            Alignment,
            Border,
            Font,
            PatternFill,
            Side,
        )
        from openpyxl.utils import get_column_letter
    except ImportError as e:
        raise ExportError("openpyxl 未安装") from e

    snapshots, total, actual_date = get_screened_stocks(
        trade_date=trade_date, min_score=0, page=1, page_size=5000
    )

    if not snapshots:
        raise ExportError(f"日期 {actual_date} 无筛选结果，无法导出")

    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None  # 新建 Workbook 必有默认 sheet
    ws.title = f"筛选结果_{actual_date}"

    # ── 样式定义 ──────────────────────────────────────────────────────────
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
    alt_fill = PatternFill(start_color="EBF3FB", end_color="EBF3FB", fill_type="solid")
    normal_font = Font(name="微软雅黑", size=10)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=False)
    thin_border = Border(
        left=Side(style="thin", color="D0D7E5"),
        right=Side(style="thin", color="D0D7E5"),
        top=Side(style="thin", color="D0D7E5"),
        bottom=Side(style="thin", color="D0D7E5"),
    )

    # ── 标题行 ────────────────────────────────────────────────────────────
    headers = [
        "股票代码", "公司名称", "行业", "收盘价", "涨跌幅%",
        "PE_TTM", "扣非PE_TTM", "市净率PB", "股息率%",
        "总市值(亿)", "综合评分", "全市场排名%",
    ]

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    # 冻结首行
    ws.freeze_panes = "A2"

    # ── 数据行 ────────────────────────────────────────────────────────────
    for row_idx, snap in enumerate(snapshots, start=2):
        row_data = [
            snap.ts_code,
            snap.name,
            snap.industry,
            round(snap.close, 2) if snap.close else "",
            round(snap.pct_chg, 2) if snap.pct_chg else "",
            round(snap.pe_ttm, 1) if snap.pe_ttm else "",
            round(snap.pe_deduct_ttm, 1) if snap.pe_deduct_ttm else "",
            round(snap.pb, 2) if snap.pb else "",
            round(snap.dv_ttm, 2) if snap.dv_ttm else "",
            round(snap.total_mv_yi, 1) if snap.total_mv_yi else "",
            round(snap.composite_score, 1) if snap.composite_score else "",
            round(snap.composite_rank, 1) if snap.composite_rank else "",
        ]

        fill = alt_fill if row_idx % 2 == 0 else None

        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = normal_font
            cell.alignment = center_align
            cell.border = thin_border
            if fill:
                cell.fill = fill

            # 涨跌幅着色
            if col_idx == 5 and isinstance(value, (int, float)):
                if value > 0:
                    cell.font = Font(name="微软雅黑", size=10, color="C00000")
                elif value < 0:
                    cell.font = Font(name="微软雅黑", size=10, color="00B050")

    # ── 列宽 ─────────────────────────────────────────────────────────────
    col_widths = [14, 18, 14, 9, 9, 9, 11, 9, 9, 12, 10, 12]
    for col_idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 24

    # ── 汇总行 ────────────────────────────────────────────────────────────
    summary_row = len(snapshots) + 2
    ws.cell(row=summary_row, column=1, value=f"共 {total} 只股票通过筛选")
    ws.cell(row=summary_row, column=1).font = Font(name="微软雅黑", bold=True, size=10)

    # ── 保存 ─────────────────────────────────────────────────────────────
    settings.exports_dir.mkdir(parents=True, exist_ok=True)
    filename = f"StockSentry_{actual_date}.xlsx"
    filepath = settings.exports_dir / filename
    wb.save(filepath)

    logger.info("Excel 导出完成", file=str(filepath), rows=len(snapshots))
    return filepath
